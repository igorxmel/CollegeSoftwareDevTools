"""Собрать отчёт по ПЗ № 8 (отладка модулей и обработка исключений) в .docx.

Листинги читаются из файлов решения, а не дублируются здесь, поэтому отчёт не может
разойтись с тем, что лежит рядом и реально компилируется.

Usage:
    python .claude/tools/build_debug_report.py
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

TASK = Path(__file__).resolve().parents[2] / "Task 8th Module Debugging and Exception Handling"
SOLUTION = TASK / "Solution"
OUT = SOLUTION / "ПЗ-8_Отладка.docx"

FONT = "Times New Roman"
MONO = "Courier New"
SIZE = Pt(14)
CODE_SIZE = Pt(10)
TABLE_SIZE = Pt(12)
INDENT = Cm(1.25)

LINES_PER_PAGE = 30
GAP_AFTER_HEADER = 3
GAP_BEFORE_SIGN = 4
GAP_BEFORE_CITY = 2

TITLE = [
    "Министерство общего и профессионального образования",
    "Ростовской области",
    "Государственное бюджетное профессиональное образовательное",
    "учреждение Ростовской области",
    "«Ростовский-на-Дону колледж связи и информатики»",
]
TOPIC = ("Тема: Отладка отдельных модулей программного проекта. "
         "Организация обработки исключений")
SIGN_LEFT = ["Выполнил: студент 4 курса", "группы ИС-42", "Холкин Егор"]
SIGN_RIGHT = ["Проверил: преподаватель", "Черевко А. Д."]

GOAL = (
    "Научиться отлаживать отдельные модули программного проекта и организовывать в них "
    "обработку исключений, а также составлять набор тестов, который проверяет каждый "
    "оператор, каждое условие и каждую ветвь алгоритма."
)

TASKS = [
    "Описать внешнюю спецификацию программы: постановку задачи, входные и выходные данные.",
    "Составить алгоритм решения задачи в виде блок-схемы.",
    "Написать программу решения задачи с обработкой исключений.",
    "Составить набор тестов для отладки программы и оформить тест-кейсы с покрытием "
    "операторов, условий и решений.",
]

VARIANT = [
    "Общая постановка. Составить алгоритм и написать программу последовательного "
    "вычисления значений заданной функции Y(X) до тех пор, пока не будет пройдена "
    "некоторая характерная точка графика функции. Значения аргумента X составляют "
    "возрастающую последовательность с шагом h. Начальное значение X0 и шаг h задаёт "
    "пользователь.",

    "Вариант 6. Вид функции: Y(X) = x² − eˣ. Задача: найти точку (точки), "
    "в которой функция равна −10.",
]

THEORY = [
    "Исключение это объект, который создаётся при ошибке во время выполнения программы и "
    "прерывает обычный ход вычислений. В Java исключения делятся на проверяемые, которые "
    "компилятор заставляет либо перехватить, либо объявить в заголовке метода, и "
    "непроверяемые, наследники RuntimeException, про которые компилятор не напоминает. "
    "Перехват выполняется блоком try, обработка блоком catch, а блок finally выполняется "
    "в любом случае, и это удобное место для освобождения ресурсов.",

    "Собственное исключение заводится наследованием от Exception. Это позволяет отделить "
    "ошибки предметной области от технических: в этой работе неположительный шаг и "
    "отсутствие характерной точки это разные ситуации, и обрабатываются они по-разному, "
    "хотя обе наследуют общий класс SearchException.",

    "Отладка это поиск и устранение ошибок в уже написанной программе. Её основной "
    "инструмент это тесты, а качество набора тестов оценивают покрытием. Покрытие "
    "операторов требует, чтобы каждый оператор выполнился хотя бы один раз. Покрытие "
    "решений строже: каждая развилка должна хотя бы раз пойти по ветви «да» и хотя бы "
    "раз по ветви «нет». Покрытие условий требует, чтобы каждое простое условие внутри "
    "составного приняло оба значения, что важно для выражений с операциями || и &&.",

    "Вещественные вычисления добавляют свои ошибки. Тип double не хранит числа точно, "
    "поэтому проверять равенство двух вещественных значений оператором == нельзя: вместо "
    "этого сравнивают модуль разности с малой величиной. Кроме того, результат может "
    "выйти за пределы типа и стать бесконечностью, а разность двух бесконечностей даёт "
    "NaN, то есть «не число». Ни то, ни другое не вызывает исключения само по себе, "
    "вычисления просто продолжаются с испорченным значением, поэтому такую проверку "
    "приходится ставить вручную.",
]

SPEC_DESC = [
    "Программа вычисляет значения функции Y(X) = x² − eˣ в точках, которые "
    "образуют возрастающую последовательность: X0, X0 + h, X0 + 2h и так далее. "
    "Характерной считается точка, в которой функция принимает значение −10. Проход "
    "прекращается, как только эта точка пройдена, то есть когда отклонение функции от "
    "искомого уровня меняет знак между двумя соседними точками.",

    "Сама точка при этом почти никогда не совпадает с узлом сетки, поэтому после "
    "обнаружения смены знака её положение уточняется половинным делением отрезка между "
    "последними двумя точками. Это даёт координату с точностью до 10⁻¹⁰ "
    "независимо от того, насколько крупный шаг задал пользователь.",
]

SPEC_IN = [
    "X0 вещественное число, начальное значение аргумента. Ограничений нет.",
    "h вещественное число, шаг изменения аргумента. Должно быть строго положительным, "
    "иначе последовательность не будет возрастающей.",
]

SPEC_OUT = [
    "протокол вычислений: номер шага, значение аргумента и значение функции в нём;",
    "координаты найденной характерной точки X и Y(X), уточнённые половинным делением, "
    "и номер шага, на котором она была пройдена;",
    "сообщение об ошибке, если шаг задан неверно, если точка не встретилась за "
    "отведённое число шагов, если значение функции вышло за пределы вещественного типа "
    "или если исходные данные закончились раньше времени.",
]

MODULES = [
    ("Main.java", "управляющий модуль: связывает ввод, поиск и вывод, перехватывает "
                  "все исключения в одном месте"),
    ("InputModule.java", "ввод вещественных чисел, нечисловой ввод обрабатывается "
                         "повторным запросом"),
    ("FunctionModule.java", "вычисление функции варианта и её отклонения от искомого "
                            "уровня, проверка результата на выход за пределы типа"),
    ("SearchModule.java", "последовательный проход по аргументу, обнаружение смены знака "
                          "и уточнение точки половинным делением"),
    ("OutputModule.java", "печать заголовка, протокола, результата и сообщений об ошибках"),
    ("SearchException.java", "общий предок исключений поиска"),
    ("InvalidStepException.java", "шаг задан неположительным числом"),
    ("PointNotFoundException.java", "точка не встретилась за отведённое число шагов"),
]

EXCEPTIONS = [
    "Обработка исключений разнесена по двум уровням. Ошибки, которые пользователь может "
    "исправить прямо сейчас, обрабатываются на месте: нечисловой ввод перехватывается "
    "внутри InputModule, программа печатает сообщение и повторяет запрос. Ошибки, после "
    "которых продолжать бессмысленно, поднимаются до Main и перехватываются там.",

    "В Main четыре блока catch. InvalidStepException и PointNotFoundException это "
    "собственные проверяемые исключения, они описывают ситуации предметной области. "
    "ArithmeticException возникает, когда значение функции перестаёт быть конечным "
    "числом. NoSuchElementException приходит из модуля ввода, когда исходные данные "
    "закончились. Блок finally закрывает поток ввода независимо от того, чем кончилась "
    "работа.",
]

RESULT_NOTE = (
    "Ниже показан протокол при X0 = 0 и h = 0,1. Длинный протокол печатается с пропуском "
    "середины, чтобы результат оставался читаемым."
)

TEST_HEADER = ["№", "Входные данные", "Что проверяется", "Результат", "Итог"]

TESTS = [
    ("Т1", "X0 = 0\nh = 0,1",
     "основной ход: характерная точка существует и находится",
     "точка пройдена на шаге 30, X = 2,918827, Y(X) = −10,000000", "пройден"),
    ("Т2", "X0 = 0\nh = 0,5",
     "крупный шаг: точка находится за меньшее число шагов",
     "точка пройдена на шаге 6, уточнённое X то же самое", "пройден"),
    ("Т3", "X0 = 0\nh = 0",
     "нулевой шаг, ветвь «нет» у проверки h > 0",
     "InvalidStepException, сообщение о неверном шаге, вычисления не начинаются",
     "пройден"),
    ("Т4", "X0 = 0\nh = −0,1",
     "отрицательный шаг, та же ветвь",
     "InvalidStepException, сообщение о неверном шаге", "пройден"),
    ("Т5", "X0 = 10\nh = 0,1",
     "переполнение: eˣ перестаёт помещаться в double",
     "ArithmeticException при x = 709,8, значение функции бесконечно", "пройден"),
    ("Т6", "X0 = abc, затем 0\nh = 0,1",
     "нечисловой ввод, перехват NumberFormatException в модуле ввода",
     "сообщение об ошибке, повторный запрос, дальше работа идёт как в Т1", "пройден"),
    ("Т7", "X0 = 10\nh = 0,01",
     "предел числа шагов: точка левее старта, функция убывает",
     "PointNotFoundException после 10000 шагов, последнее X = 110", "пройден"),
    ("Т8", "введено только X0",
     "исходные данные закончились раньше времени",
     "NoSuchElementException, сообщение «исходные данные закончились»", "пройден"),
    ("Т9", "X0 = 2,918826982306345\nh = 0,1",
     "старт точно в характерной точке, ветвь «да» у проверки | D | < ε",
     "точка найдена на шаге 0, цикл не выполняется ни разу", "пройден"),
    ("Т10", "X0 = 1e200\nh = 1",
     "неопределённость: x² и eˣ оба бесконечны, разность даёт NaN",
     "ArithmeticException, значение функции не является числом", "пройден"),
]

COVER_HEADER = ["Решение в алгоритме", "Ветвь «да»", "Ветвь «нет»"]

COVERAGE = [
    ("h > 0", "Т1, Т2", "Т3, Т4"),
    ("| D | < ε", "Т9", "Т1"),
    ("значение Y(X) конечно", "Т1", "Т5, Т10"),
    ("Dпр · D ≤ 0", "Т1, шаг 30", "Т1, шаги с 1 по 29"),
    ("k < MAX_STEPS", "Т1", "Т7"),
    ("строка преобразуется в число", "Т1", "Т6"),
    ("есть очередная строка ввода", "Т1", "Т8"),
]

COVER_NOTE = [
    "Покрытие решений обеспечено полностью: в таблице 2 у каждой развилки алгоритма "
    "заполнены обе ветви.",

    "Покрытие условий проверяется отдельно на составном условии «значение не является "
    "числом или бесконечно», которое стоит в модуле вычислений. Первое простое условие "
    "истинно в тесте Т10, второе в тесте Т5, оба ложны в Т1, так что каждое из них "
    "принимает оба значения.",

    "Покрытие операторов обеспечивается набором Т1, Т3, Т5, Т6, Т7, Т8 и Т9: вместе они "
    "выполняют каждый оператор программы, включая все четыре блока catch в управляющем "
    "модуле, обе ветви печати протокола и обе ветви уточнения половинным делением.",
]

DEBUG_NOTES = [
    "Пара вещей выяснилась только на прогоне. Тесты Т5 и Т7 задумывались как один: "
    "взять старт правее характерной точки и упереться в предел числа шагов. На деле при "
    "X0 = 10 и шаге 0,1 программа до предела не доходит, потому что раньше, уже при "
    "x ≈ 709,8, экспонента перестаёт помещаться в double. Чтобы добраться до ветви с "
    "пределом шагов, шаг пришлось уменьшить до 0,01: тогда за 10000 шагов аргумент "
    "дорастает только до 110, и значение функции остаётся конечным.",

    "Проверять результат вычисления на бесконечность оказалось мало. При очень большом "
    "аргументе и x², и eˣ по отдельности обращаются в бесконечность, а их разность даёт "
    "NaN, который условием на бесконечность не ловится. Поэтому в модуле вычислений "
    "стоит составное условие из двух проверок, и тест Т10 существует именно ради второй "
    "половины этого условия.",

    "Равенство вещественных чисел нигде не проверяется оператором ==. Попадание точно в "
    "характерную точку определяется сравнением модуля отклонения с малой величиной ε, а "
    "сам факт того, что точка пройдена, по смене знака произведения Dпр · D. Иначе "
    "условие не сработало бы почти никогда: X = 2,918827 это бесконечная дробь, и "
    "попасть в неё узлом сетки практически невозможно.",
]

CONCLUSION = [
    "В работе я разложил программу на модули ввода, вычисления, поиска и вывода и "
    "организовал обработку исключений на двух уровнях: исправимые ошибки обрабатываются "
    "там же, где возникают, а неисправимые поднимаются до управляющего модуля. Отдельно "
    "завёл два собственных исключения, потому что неверный шаг и ненайденная точка это "
    "разные ситуации, и реакция на них должна различаться.",

    "Набор из десяти тестов закрывает все ветви алгоритма и оба значения каждого простого "
    "условия. Полезнее всего оказались тесты на границы вещественного типа: ошибка с NaN "
    "не вызывает исключения сама по себе, программа продолжила бы считать с испорченным "
    "значением и выдала бы неверный ответ без единого признака сбоя. Без теста Т10 я бы "
    "её не заметил.",
]


def setup(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = SIZE
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = INDENT
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def para(doc, text="", *, indent=INDENT, align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = indent
    p.paragraph_format.alignment = align
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    if text:
        run = p.add_run(text)
        run.font.name = FONT
        run.font.size = SIZE
        run.bold = bold
    return p


def code(doc, text):
    for line in text.splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(line)
        run.font.name = MONO
        run.font.size = CODE_SIZE
        run.element.rPr.rFonts.set(qn("w:eastAsia"), MONO)


def blank(doc, n=1):
    for _ in range(n):
        para(doc)


def heading(doc, number, title, *, new_page=True):
    if new_page:
        para(doc).add_run().add_break(WD_BREAK.PAGE)
    para(doc, f"{number} {title}", align=WD_ALIGN_PARAGRAPH.LEFT, bold=True)
    blank(doc)


def table(doc, header, rows, widths):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False

    def fill(cells, values, bold):
        for cell, value, width in zip(cells, values, widths):
            cell.width = Cm(width)
            cell.text = ""
            for i, line in enumerate(str(value).split("\n")):
                p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if bold else WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(line)
                run.font.name = FONT
                run.font.size = TABLE_SIZE
                run.bold = bold

    fill(t.rows[0].cells, header, True)
    for row in rows:
        fill(t.add_row().cells, row, False)
    return t


def caption(doc, text):
    para(doc, text, indent=Cm(0), align=WD_ALIGN_PARAGRAPH.CENTER)


def title_page(doc) -> int:
    lines = 0
    for line in TITLE:
        para(doc, line, indent=Cm(0), align=WD_ALIGN_PARAGRAPH.CENTER)
    lines += len(TITLE)
    blank(doc, GAP_AFTER_HEADER)
    lines += GAP_AFTER_HEADER
    para(doc, "ОТЧЁТ", indent=Cm(0), align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    para(doc, "по практическому занятию № 8", indent=Cm(0),
         align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    blank(doc)
    para(doc, TOPIC, indent=Cm(0), align=WD_ALIGN_PARAGRAPH.CENTER)
    lines += 2 + 1 + max(1, -(-len(TOPIC) // 67))
    blank(doc, GAP_BEFORE_SIGN)
    lines += GAP_BEFORE_SIGN
    for line in SIGN_LEFT:
        para(doc, line, indent=Cm(0), align=WD_ALIGN_PARAGRAPH.RIGHT)
    blank(doc)
    for line in SIGN_RIGHT:
        para(doc, line, indent=Cm(0), align=WD_ALIGN_PARAGRAPH.RIGHT)
    lines += len(SIGN_LEFT) + 1 + len(SIGN_RIGHT)
    blank(doc, GAP_BEFORE_CITY)
    lines += GAP_BEFORE_CITY
    para(doc, "Ростов-на-Дону", indent=Cm(0), align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "2026", indent=Cm(0), align=WD_ALIGN_PARAGRAPH.CENTER)
    return lines + 2


def build():
    doc = Document()
    setup(doc)

    used = title_page(doc)
    if used > LINES_PER_PAGE:
        raise SystemExit(f"титульный лист занял {used} строк при лимите {LINES_PER_PAGE}")

    heading(doc, "1", "Цель работы")
    para(doc, GOAL)

    heading(doc, "2", "Задание")
    for i, text in enumerate(TASKS, start=1):
        para(doc, f"{i}. {text}")
    blank(doc)
    for text in VARIANT:
        para(doc, text)

    heading(doc, "3", "Краткие теоретические сведения")
    for text in THEORY:
        para(doc, text)

    heading(doc, "4", "Внешняя спецификация")
    para(doc, "4.1 Описание задачи", align=WD_ALIGN_PARAGRAPH.LEFT, bold=True)
    blank(doc)
    for text in SPEC_DESC:
        para(doc, text)
    blank(doc)
    para(doc, "4.2 Входные данные", align=WD_ALIGN_PARAGRAPH.LEFT, bold=True)
    blank(doc)
    for text in SPEC_IN:
        para(doc, f"- {text}")
    blank(doc)
    para(doc, "4.3 Выходные данные", align=WD_ALIGN_PARAGRAPH.LEFT, bold=True)
    blank(doc)
    for text in SPEC_OUT:
        para(doc, f"- {text}")

    heading(doc, "5", "Блок-схема алгоритма")
    para(doc, "Алгоритм решения задачи приведён на рисунке 1.")
    blank(doc)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run().add_picture(str(SOLUTION / "flowchart.png"), width=Cm(14.5))
    caption(doc, "Рисунок 1 - Блок-схема алгоритма")

    heading(doc, "6", "Код программы")
    para(doc, "Программа написана на языке Java и разбита на модули:")
    blank(doc)
    for name, role in MODULES:
        para(doc, f"- {name} - {role};")
    blank(doc)
    for text in EXCEPTIONS:
        para(doc, text)
    blank(doc)
    for name, _ in MODULES:
        para(doc, f"Листинг файла {name}:", indent=Cm(0))
        code(doc, (SOLUTION / name).read_text(encoding="utf-8").rstrip())
        blank(doc)

    heading(doc, "7", "Тестирование и отладка")
    para(doc, RESULT_NOTE)
    blank(doc)
    code(doc, (SOLUTION / "run-example.txt").read_text(encoding="utf-8").rstrip())
    blank(doc)
    caption(doc, "Таблица 1 - Тест-кейсы")
    table(doc, TEST_HEADER, TESTS, [1.0, 3.4, 4.5, 5.0, 2.1])
    blank(doc)
    caption(doc, "Таблица 2 - Покрытие решений")
    table(doc, COVER_HEADER, COVERAGE, [6.8, 4.6, 4.6])
    blank(doc)
    for text in COVER_NOTE:
        para(doc, text)
    blank(doc)
    for text in DEBUG_NOTES:
        para(doc, text)

    heading(doc, "8", "Вывод")
    for text in CONCLUSION:
        para(doc, text)

    try:
        doc.save(OUT)
    except PermissionError:
        raise SystemExit(f"{OUT.name} открыт в Word, закрой его и повтори")
    print(f"saved: {OUT.name}")


if __name__ == "__main__":
    build()
