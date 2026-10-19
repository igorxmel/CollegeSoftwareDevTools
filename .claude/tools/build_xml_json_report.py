"""Собрать отчёт по ПЗ № 7 (XML и JSON) в .docx.

Листинги не дублируются в скрипте: они читаются из файлов решения, поэтому отчёт не
может разойтись с тем, что лежит рядом и реально работает.

Usage:
    python .claude/tools/build_xml_json_report.py
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

TASK = Path(__file__).resolve().parents[2] / "Task 7th XML and JSON"
SOLUTION = TASK / "Solution"
OUT = SOLUTION / "ПЗ-7_XML_JSON.docx"

FONT = "Times New Roman"
MONO = "Courier New"
SIZE = Pt(14)
CODE_SIZE = Pt(10)
INDENT = Cm(1.25)

LINES_PER_PAGE = 30
GAP_AFTER_HEADER = 3
GAP_BEFORE_SIGN = 4
GAP_BEFORE_CITY = 2

TITLE = [
    "Министерство общего и профессионального образования",
    "Ростовской области",
    "Государственное бюджетное профессиональное образовательное",
    "учреждение Ростовской области",
    "«Ростовский-на-Дону колледж связи и информатики»",
]
TOPIC = "Тема: Представление информации в XML и JSON"
SIGN_LEFT = ["Выполнил: студент 4 курса", "группы ИС-42", "Холкин Егор"]
SIGN_RIGHT = ["Проверил: преподаватель", "Черевко А. Д."]

GOAL = (
    "Научиться описывать данные в форматах XML и JSON, понять разницу между ними, а также "
    "освоить преобразование JSON в объект JavaScript и вывод его содержимого на "
    "HTML-страницу средствами DOM."
)

TASKS = [
    "Представить сведения о книге в виде корректного XML-документа.",
    "Сформировать XML-документ с данными абонента телефонной линии.",
    "Описать данные в формате JSON, преобразовать их в объект JavaScript и отобразить "
    "содержимое объекта в элементе HTML-страницы.",
    "Выполнить парсинг строкового JSON в объект JavaScript для дальнейшей работы с "
    "данными.",
]

THEORY = [
    "XML это расширяемый язык разметки. Он описывает данные при помощи вложенных тегов, "
    "имена которых придумывает сам разработчик. Документ обязан быть корректным: ровно "
    "один корневой элемент, каждый открытый тег закрыт, вложенность не нарушена, "
    "значения атрибутов в кавычках. В первой строке обычно стоит объявление с версией и "
    "кодировкой.",

    "Значения в XML можно хранить двумя способами: содержимым элемента или атрибутом. "
    "Атрибуты удобны для коротких служебных характеристик, элементы для самих данных. В "
    "документе об абоненте я вынес серию и номер паспорта в атрибуты, потому что это "
    "части одного реквизита, а адрес расписал вложенными элементами, так как каждую его "
    "часть может понадобиться прочитать отдельно.",

    "JSON это текстовый формат обмена данными, построенный на синтаксисе объектов "
    "JavaScript, но независимый от языка. Данные записываются парами «ключ и значение» в "
    "фигурных скобках, списки в квадратных. Ключи и строковые значения всегда в двойных "
    "кавычках, числа и логические значения без них, завершающая запятая недопустима. "
    "Типы значений: строка, число, логическое значение, null, объект, массив.",

    "Строку JSON превращает в объект метод JSON.parse, обратное преобразование делает "
    "JSON.stringify. После разбора это обычный объект, к полям которого обращаются через "
    "точку, а вывести содержимое на страницу можно через свойства textContent или "
    "innerHTML найденного элемента.",
]

WORK = [
    ("4.1", "Задание 1. XML-документ о книге",
     "Сведения о книге описаны вложенными элементами. Автор разбит на фамилию и "
     "инициалы: так к ним можно обратиться по отдельности, например отсортировать "
     "каталог по фамилии.",
     "book.xml"),
    ("4.2", "Задание 2. XML-документ с данными абонента",
     "Серия и номер паспорта записаны атрибутами одного элемента, потому что вместе они "
     "образуют один реквизит. Адрес, наоборот, разложен на город, улицу, дом и квартиру.",
     "subscriber.xml"),
    ("4.3", "Задания 3 и 4. JSON, преобразование в объект и парсинг строки",
     "Обе части сделаны на одной странице. В третьем задании те же сведения о книге "
     "записаны в JSON, разобраны методом JSON.parse и выведены в элемент pre перебором "
     "ключей объекта. В четвёртом разбирается строка с именем и возрастом, после чего к "
     "её полям идёт обращение через точку, выполняется арифметика над возрастом и "
     "обратное преобразование в строку.",
     "index.html"),
]

RESULT = [
    "Страница открывается в браузере без ошибок в консоли. В блоке третьего задания "
    "выводится пять строк со сведениями о книге, в блоке четвёртого тип значения после "
    "парсинга (object), имя, возраст, возраст плюс год и результат обратного "
    "преобразования в строку.",
]

ANSWERS = [
    ("Каково назначение XML?",
     "XML предназначен для описания и хранения структурированных данных в виде, понятном "
     "и человеку, и программе. Он не выводит данные на экран и не задаёт их оформление, а "
     "только описывает структуру и смысл, поэтому используется для обмена данными между "
     "системами, конфигурационных файлов и документов со сложной вложенностью."),
    ("Каково назначение JSON?",
     "JSON решает ту же задачу обмена данными, но компактнее. Он построен на синтаксисе "
     "объектов JavaScript, поэтому в веб-приложениях данные с сервера приходят обычно "
     "именно в нём: разбор занимает одну строку кода и сразу даёт готовый объект."),
    ("В чём отличие XML и JSON?",
     "XML описывает данные тегами и поддерживает атрибуты, пространства имён, "
     "комментарии и проверку по схеме. JSON описывает данные парами «ключ и значение», "
     "атрибутов и комментариев у него нет. XML многословнее, зато выразительнее; JSON "
     "короче и ближе к структурам данных языков программирования."),
    ("В чём состоит преимущество JSON перед XML?",
     "Он занимает меньше места, потому что имя поля пишется один раз, а не в открывающем "
     "и закрывающем теге. Его разбор проще и быстрее, а результатом сразу становится "
     "объект языка, тогда как XML приходится обходить через DOM. Синтаксис минимальный, "
     "поэтому его быстрее читать и труднее сломать."),
    ("В чём отличие XML от HTML?",
     "У HTML фиксированный набор тегов, и назначение его в том, чтобы показать документ в "
     "браузере. В XML имена тегов придумывает разработчик, и назначение его в том, чтобы "
     "описать данные, а не их вид. HTML прощает ошибки разметки, браузер отобразит "
     "страницу и с незакрытым тегом, XML же обязан быть корректным, иначе анализатор "
     "откажется его читать."),
]

CONCLUSION = [
    "В работе я описал одни и те же сведения двумя способами и увидел разницу на "
    "практике. XML заставляет писать больше: каждое значение обрамляется парой тегов, "
    "зато появляется место для атрибутов и комментариев. JSON короче и сразу ложится в "
    "объект языка, но описать в нём что-то вроде серии и номера паспорта отдельными "
    "атрибутами уже не получится, придётся заводить вложенный объект.",

    "Из этого следует практический вывод: для обмена данными между страницей и сервером "
    "удобнее JSON, а XML имеет смысл там, где важны строгая проверка структуры по схеме "
    "и служебные характеристики рядом со значениями.",
]


def setup(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = SIZE
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = INDENT
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def para(doc, text="", *, indent=INDENT, align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = indent
    p.paragraph_format.alignment = align
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    if text:
        run = p.add_run(text)
        run.font.name = FONT
        run.font.size = SIZE
        run.bold = bold
    return p


def code(doc, text):
    for line in text.splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(line)
        run.font.name = MONO
        run.font.size = CODE_SIZE
        run.element.rPr.rFonts.set(qn("w:eastAsia"), MONO)


def blank(doc, n=1):
    for _ in range(n):
        para(doc)


def heading(doc, number, title, *, new_page=True):
    if new_page:
        para(doc).add_run().add_break(WD_BREAK.PAGE)
    para(doc, f"{number} {title}", align=WD_ALIGN_PARAGRAPH.LEFT, bold=True)
    blank(doc)


def title_page(doc) -> int:
    lines = 0
    for line in TITLE:
        para(doc, line, indent=Cm(0), align=WD_ALIGN_PARAGRAPH.CENTER)
    lines += len(TITLE)
    blank(doc, GAP_AFTER_HEADER)
    lines += GAP_AFTER_HEADER
    para(doc, "ОТЧЁТ", indent=Cm(0), align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    para(doc, "по практическому занятию № 7", indent=Cm(0),
         align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    blank(doc)
    para(doc, TOPIC, indent=Cm(0), align=WD_ALIGN_PARAGRAPH.CENTER)
    lines += 2 + 1 + max(1, -(-len(TOPIC) // 67))
    blank(doc, GAP_BEFORE_SIGN)
    lines += GAP_BEFORE_SIGN
    for line in SIGN_LEFT:
        para(doc, line, indent=Cm(0), align=WD_ALIGN_PARAGRAPH.RIGHT)
    blank(doc)
    for line in SIGN_RIGHT:
        para(doc, line, indent=Cm(0), align=WD_ALIGN_PARAGRAPH.RIGHT)
    lines += len(SIGN_LEFT) + 1 + len(SIGN_RIGHT)
    blank(doc, GAP_BEFORE_CITY)
    lines += GAP_BEFORE_CITY
    para(doc, "Ростов-на-Дону", indent=Cm(0), align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "2026", indent=Cm(0), align=WD_ALIGN_PARAGRAPH.CENTER)
    return lines + 2


def build():
    doc = Document()
    setup(doc)

    used = title_page(doc)
    if used > LINES_PER_PAGE:
        raise SystemExit(f"титульный лист занял {used} строк при лимите {LINES_PER_PAGE}")

    heading(doc, "1", "Цель работы")
    para(doc, GOAL)

    heading(doc, "2", "Задание")
    for i, text in enumerate(TASKS, start=1):
        para(doc, f"{i}. {text}")

    heading(doc, "3", "Краткие теоретические сведения")
    for text in THEORY:
        para(doc, text)

    heading(doc, "4", "Ход работы")
    for number, title, comment, filename in WORK:
        para(doc, f"{number} {title}", align=WD_ALIGN_PARAGRAPH.LEFT, bold=True)
        blank(doc)
        para(doc, comment)
        blank(doc)
        para(doc, f"Листинг файла {filename}:", indent=Cm(0))
        code(doc, (SOLUTION / filename).read_text(encoding="utf-8").rstrip())
        blank(doc)

    para(doc, "4.4 Результат работы страницы", align=WD_ALIGN_PARAGRAPH.LEFT, bold=True)
    blank(doc)
    for text in RESULT:
        para(doc, text)

    heading(doc, "5", "Ответы на контрольные вопросы")
    for i, (question, answer) in enumerate(ANSWERS, start=1):
        para(doc, f"{i} {question}")
        para(doc, answer)
        blank(doc)

    heading(doc, "6", "Вывод")
    for text in CONCLUSION:
        para(doc, text)

    try:
        doc.save(OUT)
    except PermissionError:
        raise SystemExit(f"{OUT.name} открыт в Word, закрой его и повтори")
    print(f"saved: {OUT.name}")


if __name__ == "__main__":
    build()
